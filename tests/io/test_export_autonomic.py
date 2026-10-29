"""
Autonomic (CSI / CPI) results in the exports: the Excel Autonomic /
Autonomic_Series / Autonomic_LFHF / EventLocked sheets, the Cover rows
stating the method, window, ks / kp and the rat-calibration caveat, the
Prism tables and project notes, and the PDF header line.
"""
import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from ecg.core.autonomic import beat_dataframe, compute_autonomic_indices, with_coefficients
from ecg.io.export import (
    ExcelExporter, PrismExporter, autonomic_calibration_caveat,
    autonomic_disclosure_rows, autonomic_pdf_line,
)

from tests.helpers import rpeaks_from_rr

FS = 2000.0


@pytest.fixture(scope="module")
def result():
    rng = np.random.default_rng(2)
    rr = 120 + rng.normal(0, 1.5, 1500)
    return compute_autonomic_indices(rpeaks_from_rr(list(rr), FS), FS)


def _spwvd(res):
    t = np.arange(float(res.t[0]), float(res.t[-1]), 0.05)
    return {"t": t, "lf": 100 + np.sin(t), "hf": 50 + np.cos(t),
            "params": {"lf_band": (0.4, 1.5), "hf_band": (1.5, 5.0),
                       "fs_resample": 20.0, "fbins": 4096}}


def _events():
    return pd.DataFrame({
        "label": ["dobu", "saline", ""], "onset_s": [60.0, 100.0, 1.0],
        "valid": [True, True, False], "coverage_pre": [1.0, 1.0, 0.1],
        "coverage_post": [1.0, 1.0, 1.0],
        "csi_pre": [10.0, 11.0, np.nan], "csi_post": [12.0, 11.5, np.nan],
        "csi_delta": [2.0, 0.5, np.nan],
        "cpi_pre": [20.0, 21.0, np.nan], "cpi_post": [18.0, 20.0, np.nan],
        "cpi_delta": [-2.0, -1.0, np.nan],
    })


def _cover(wb):
    ws = wb["Cover"]
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1)}


def _wb():
    wb = Workbook()
    wb.active.title = "Cover"
    wb["Cover"]["A1"] = "ECG Analysis Report"
    return wb


# ── disclosure ────────────────────────────────────────────────────────────

def test_disclosure_states_method_window_gains_and_the_rat_caveat(result):
    rows = dict(autonomic_disclosure_rows(result.params))
    assert "doi:10.1113/JP288400" in rows["Autonomic indices"]
    assert rows["Autonomic method"].startswith("Robust shrinkage") and "[robust]" in rows["Autonomic method"]
    assert rows["Autonomic window"].startswith("3 s sliding Poincaré window (≥ 10")
    assert rows["Autonomic ks / kp"] == "ks = 3, kp = 12"
    assert "rat (SHR)" in rows["ks / kp calibration"] and "NOT validated in mice" in rows["ks / kp calibration"]
    user = dict(autonomic_disclosure_rows(with_coefficients(result, 4.5, 9).params))
    assert user["Autonomic ks / kp"] == "ks = 4.5, kp = 9"
    assert user["ks / kp calibration"].startswith("ks = 4.5, kp = 9: user-set")
    assert "not validated in mice" in user["ks / kp calibration"]
    more = dict(autonomic_disclosure_rows(
        result.params, _spwvd(result)["params"],
        {"n_events": 3, "n_valid": 2, "pre_s": 30.0, "post_s": 20.0, "label_filter": "All"}))
    assert more["LF / HF (SPWVD)"].startswith("LF 0.4–1.5 Hz, HF 1.5–5 Hz")
    assert more["Event-locked"].startswith("3 event(s), 2 valid · pre 30 s / post 20 s")
    assert autonomic_calibration_caveat(None, 12).startswith("ks = nan")


def test_a_withheld_result_says_why():
    empty = compute_autonomic_indices(rpeaks_from_rr([120.0] * 10, FS), FS)
    rows = dict(autonomic_disclosure_rows(empty.params))
    assert rows["Autonomic withheld"].startswith("nothing computed — recording too short")


def test_pdf_line(result):
    line = autonomic_pdf_line(result.params, 4)
    assert "window 3 s" in line and "ks = 3, kp = 12" in line
    assert "rat (SHR) calibration, NOT validated in mice" in line and "4 event(s)" in line
    assert "user-set gains" in autonomic_pdf_line(with_coefficients(result, 2, 2).params)


# ── Excel ─────────────────────────────────────────────────────────────────

def test_excel_sheets_cover_rows_and_values(result):
    wb = _wb()
    added = ExcelExporter.add_autonomic_sheets(
        wb, result, _spwvd(result), _events(),
        {"n_events": 3, "n_valid": 2, "pre_s": 30.0, "post_s": 30.0, "label_filter": "All",
         "notes": ["Within-recording statistics: not population inference."]})
    assert added == ["Autonomic", "Autonomic_Series", "Autonomic_LFHF", "EventLocked"]
    cover = _cover(wb)
    for key in ("Autonomic indices", "Autonomic method", "Autonomic window",
                "Autonomic ks / kp", "ks / kp calibration", "LF / HF (SPWVD)", "Event-locked"):
        assert key in cover, key
    assert "NOT validated in mice" in cover["ks / kp calibration"]

    ws = wb["Autonomic"]
    beat = beat_dataframe(result)
    header = [ws.cell(row=1, column=c).value for c in range(1, len(beat.columns) + 1)]
    assert header == list(beat.columns)
    assert ws.max_row >= len(beat) + 1
    i = int(np.flatnonzero(np.isfinite(beat["SD1_ms"].to_numpy()))[0])
    assert ws.cell(row=i + 2, column=3).value == pytest.approx(beat["SD1_ms"].iloc[i], rel=1e-15)
    block = {ws.cell(row=r, column=len(beat.columns) + 2).value:
             ws.cell(row=r, column=len(beat.columns) + 3).value for r in range(1, ws.max_row + 1)}
    assert block["SD01_ms (whole-recording SD1)"] == pytest.approx(result.sd01)
    assert block["method"] == "robust" and block["ks"] == 3.0
    assert "LF_mean_au" in block and "ks / kp calibration" in block

    ws_s = wb["Autonomic_Series"]
    assert [ws_s.cell(row=1, column=c).value for c in range(1, 4)] == ["Time_s", "HR_bpm", "CSI"]
    assert ws_s.max_row == len(result.t) + 1
    ws_e = wb["EventLocked"]
    assert [ws_e.cell(row=1, column=c).value for c in range(1, 4)] == ["label", "onset_s", "valid"]
    assert ws_e.cell(row=4, column=3).value is False
    values = [ws_e.cell(row=r, column=2).value for r in range(1, ws_e.max_row + 1)]
    assert "Within-recording statistics: not population inference." in values


def test_excel_long_series_are_decimated_with_a_note_and_no_lfhf_without_spwvd(result):
    wb = _wb()
    added = ExcelExporter.add_autonomic_sheets(wb, result, {"error": "x"}, None, None,
                                               max_series_rows=1000)
    assert added == ["Autonomic", "Autonomic_Series"]
    ws_s = wb["Autonomic_Series"]
    assert ws_s.max_row <= 1001
    notes = [ws_s.cell(row=2, column=c).value for c in range(1, 14)]
    assert any(isinstance(v, str) and "Decimated 1:" in v for v in notes)
    assert "Event-locked" not in _cover(wb)


def test_excel_without_a_cover_sheet_still_writes_the_sheets(result):
    wb = Workbook()
    wb.active.title = "Other"
    assert "Autonomic" in ExcelExporter.add_autonomic_sheets(wb, result)


# ── Prism ─────────────────────────────────────────────────────────────────

def test_prism_tables_and_notes(result, tmp_path):
    pytest.importorskip("pzfx")
    path = tmp_path / "out.pzfx"
    results = {"hr": {"n": 3, "mean": 500.0, "min": 480.0, "max": 520.0, "std": 5.0}}
    n = PrismExporter.build_and_write(
        path, results, subject="m1", autonomic=result, autonomic_spwvd=_spwvd(result),
        event_table=_events(),
        event_params={"n_events": 3, "n_valid": 2, "pre_s": 30.0, "post_s": 30.0})
    xml = path.read_text(encoding="utf-8")
    for name in ("Autonomic — CSI / CPI", "Autonomic — D / SD1 / SD2",
                 "Autonomic — LF / HF (SPWVD)", "Event-locked — pre / post (valid events)"):
        assert name in xml, name
    assert n == 4
    for needle in ("Autonomic method", "ks / kp calibration", "NOT validated in mice",
                   "dobu @ 60.0 s", "CSI pre", "CPI post"):
        assert needle in xml, needle
    assert "@ 1.0 s" not in xml                                 # the invalid event


def test_prism_without_autonomic_is_unchanged(tmp_path):
    pytest.importorskip("pzfx")
    path = tmp_path / "out.pzfx"
    results = {"hr": {"n": 3, "mean": 500.0, "min": 480.0, "max": 520.0, "std": 5.0},
               "hrv_time": pd.DataFrame({"HRV_SDNN": [5.0]})}
    PrismExporter.build_and_write(path, results, subject="m1")
    xml = path.read_text(encoding="utf-8")
    assert "Autonomic" not in xml and "ks / kp" not in xml


# ── review fixes: the span the indices cover; Autonomic-only export ───────

def test_the_span_of_the_indices_is_disclosed_against_the_core_window(result):
    from ecg.io.export import autonomic_span_text
    p = dict(result.params, analysis_window=(0.0, 0.0), peak_span_s=(0.12, 180.4))
    assert autonomic_span_text(p) == "R peaks 0.1 – 180.4 s · analysis window full signal"
    assert autonomic_span_text(p, (0.0, 0.0)).endswith("(the same as the core HRV results)")
    other = autonomic_span_text(p, (100.0, 200.0))
    assert "NOT the core HRV results' window (100 – 200 s)" in other
    rows = dict(autonomic_disclosure_rows(p, hrv_window=(100.0, 0.0)))
    assert "100 s – end" in rows["Autonomic span"]
    # an older result without the span: no row
    assert "Autonomic span" not in dict(autonomic_disclosure_rows(result.params))
    assert autonomic_span_text({}) == ""
    line = autonomic_pdf_line(dict(p, analysis_window=(100.0, 200.0)))
    assert "analysis window 100 – 200 s" in line
    assert "analysis window" not in autonomic_pdf_line(p)          # full signal: nothing to say
    wb = _wb()
    ExcelExporter.add_autonomic_sheets(wb, _with_params(result, p), hrv_window=(100.0, 200.0))
    assert "NOT the core HRV" in _cover(wb)["Autonomic span"]


def _with_params(result, params):
    import dataclasses
    return dataclasses.replace(result, params=dict(params))


def test_autonomic_only_workbook_needs_no_core_results(result):
    wb = ExcelExporter.build_autonomic_workbook("M7", "C:/data/mouse_7.mat", 2000.0,
                                                np.array([0.0, 240.0]))
    added = ExcelExporter.add_autonomic_sheets(wb, result, _spwvd(result))
    assert wb.sheetnames == ["Cover", "Autonomic", "Autonomic_Series", "Autonomic_LFHF"]
    assert added == wb.sheetnames[1:]
    cover = _cover(wb)
    assert cover["Subject"] == "M7" and cover["File"] == "mouse_7.mat"
    assert cover["Sampling rate"] == "2000 Hz" and cover["Duration"] == "240.0 s"
    assert "Analyze was not run" in cover["Core HRV"]
    assert "ks / kp calibration" in cover
    assert "Autonomic only" in wb["Cover"]["A1"].value
    ws_l = wb["Autonomic_LFHF"]
    notes = [c.value for row in ws_l.iter_rows(min_row=1, max_row=3, min_col=4, max_col=8)
             for c in row if isinstance(c.value, str) and "a.u." in c.value]
    assert len(notes) == 1 and "its own series" in notes[0] and "bridged" not in notes[0]


def test_export_controller_writes_the_autonomic_results_before_analyze(result, tmp_path,
                                                                        monkeypatch):
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from ecg.ui import export_controller as ec
    from ecg.ui.state import AnalysisState, SignalState

    warnings_shown = []
    monkeypatch.setattr(ec.messagebox, "showwarning",
                        lambda title, msg: warnings_shown.append(msg))
    statuses = []
    analysis = AnalysisState()
    analysis.autonomic_result = result
    app = SimpleNamespace(
        analysis=analysis,
        signal=SignalState(fs=2000.0, filepath="C:/data/mouse_7.mat",
                           time=np.array([0.0, 200.0])),
        ent_subject=SimpleNamespace(get=lambda: "M7"),
        _set_status=lambda text, color=None: statuses.append((text, color)))
    ctrl = ec.ExportController(app)
    assert ctrl.autonomic_only() and ctrl.hrv_window() is None
    path = tmp_path / "auto.xlsx"
    monkeypatch.setattr(ec.filedialog, "asksaveasfilename", lambda **kw: str(path))
    ctrl.export_excel()
    assert path.exists() and "Autonomic only" in statuses[-1][0]
    wb = load_workbook(path)
    assert {"Cover", "Autonomic", "Autonomic_Series"} <= set(wb.sheetnames)
    assert "HR_Statistics" not in wb.sheetnames
    # the exports that need core results say where the Autonomic ones go
    ctrl.export_zip()
    ctrl.export_prism()
    assert len(warnings_shown) == 2
    assert all("Export Excel saves them alone" in w for w in warnings_shown)
    # without any result: the plain warning
    analysis.autonomic_result = None
    ctrl.export_excel()
    assert warnings_shown[-1] == "Run analysis first."
