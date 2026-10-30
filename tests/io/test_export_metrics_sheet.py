"""
The Metrics sheet is additive: one line per metric-table row with raw values,
and every other sheet is written exactly as without it.
"""
import io
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from ecg.core import analysis
from ecg.core.analysis import analyse_core
from ecg.core.metric_table import build_metric_table
from ecg.io.export import ExcelExporter
from ecg.ui.export_controller import ExportController, pdf_metric_rows
from ecg.ui.state import AnalysisState, DetectionState, SignalState

from tests.helpers import make_synth_ecg

pytestmark = pytest.mark.skipif(not analysis.NK_AVAILABLE, reason="neurokit2 not installed")


@pytest.fixture(scope="module")
def recording():
    sig, fs, rp = make_synth_ecg(n_beats=80, seed=4)
    return sig, fs, rp, analyse_core(sig, rp, fs)


def _sheet_rows(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _app(sig, fs, rp, r, **analysis_kw):
    return SimpleNamespace(
        analysis=AnalysisState(results=r, analysed_window=(0.0, 0.0), **analysis_kw),
        detection=DetectionState(rpeaks_ok=rp, sig_quality=80),
        signal=SignalState(fs=fs, filtered=sig, raw_norm=sig, time=np.arange(len(sig)) / fs,
                           filepath="C:/data/m1.mat"),
        ent_subject=SimpleNamespace(get=lambda: "m1"),
        _qtc_formula=lambda: "mitchell",
    )


def _read(buf):
    buf.seek(0)
    return load_workbook(buf)


def test_sheet_holds_raw_values_and_empty_cells_for_missing(recording):
    _, _, _, r = recording
    wb = Workbook()
    table = build_metric_table(r, states={"core": "up to date"})
    ExcelExporter.add_metrics_sheet(wb, table)
    buf = io.BytesIO()
    wb.save(buf)
    ws = _read(buf)["Metrics"]
    rows = _sheet_rows(ws)
    header = rows[0]
    assert header == ["Key", "Module", "Group", "Metric", "Value", "Unit", "Basis_n",
                      "Basis_total", "Basis", "State", "Tier", "Note"]
    by_key = {row[0]: dict(zip(header, row)) for row in rows[1:]}
    assert len(by_key) == len(table)
    assert by_key["hr.mean"]["Value"] == r["hr"]["mean"]     # not rounded
    assert by_key["hr.mean"]["State"] == "up to date"
    assert by_key["hr.mean"]["Basis_n"] == r["hr"]["n_valid"]
    assert by_key["hrv_time.HRV_SDNN"]["Value"] == float(r["hrv_time"]["HRV_SDNN"].values[0])
    assert by_key["epochs.status"]["Value"] is None
    assert by_key["epochs.status"]["Basis_n"] is None


def test_write_excel_adds_only_the_metrics_sheet(recording, monkeypatch):
    sig, fs, rp, r = recording
    app = _app(sig, fs, rp, r)
    buf = io.BytesIO()
    ExportController(app).write_excel(buf)
    with_sheet = _read(buf)

    monkeypatch.setattr(ExcelExporter, "add_metrics_sheet", staticmethod(lambda wb, table: None))
    buf2 = io.BytesIO()
    ExportController(app).write_excel(buf2)
    without = _read(buf2)

    assert with_sheet.sheetnames == without.sheetnames + ["Metrics"]
    for name in without.sheetnames:
        a, b = _sheet_rows(with_sheet[name]), _sheet_rows(without[name])
        if name == "Cover":   # the analysis-date row carries the clock
            a = [row for row in a if row[0] != "Analysis date"]
            b = [row for row in b if row[0] != "Analysis date"]
        assert a == b, name


def test_a_stale_classification_is_left_out_of_the_sheet(recording):
    sig, fs, rp, r = recording
    app = _app(sig, fs, rp, r, premature_summary={"n_premature": 9, "span_s": 10.0})
    app.analysis_ctrl = SimpleNamespace(arrhythmia_is_current=lambda: False)
    t = ExportController(app).metric_table()
    assert not t.has("events.n_premature")
    app.analysis_ctrl = SimpleNamespace(arrhythmia_is_current=lambda: True)
    assert ExportController(app).metric_table().value("events.n_premature") == 9


# ── PDF rows ───────────────────────────────────────────────────────────────

def _pdf(r):
    return {label: val for label, val, _ in pdf_metric_rows(r)}


def test_pdf_rows_format_the_table_values(recording):
    _, _, _, r = recording
    rows = _pdf(r)
    assert rows["Mean HR"] == f"{r['hr']['mean']:.1f} bpm"
    assert rows["SDNN"] == f"{float(r['hrv_time']['HRV_SDNN'].values[0]):.1f} ms"
    assert rows["pNN6"] == f"{float(r['hrv_time']['HRV_pNN6'].values[0]):.1f} %"
    assert "SD1" not in rows and "PR" not in rows


def test_pdf_rows_withheld_and_legacy_spectrum():
    fd = pd.DataFrame()
    fd.attrs["withheld_reason"] = "too short"
    assert _pdf({"hrv_freq": fd})["LF / HF"] == "withheld"
    assert _pdf({"hrv_freq": pd.DataFrame({"HRV_LFn": [0.4]})})["LF / HF"].startswith("legacy")


def test_pdf_rows_lf_hf_as_percent_and_zero_left_out():
    fd = pd.DataFrame({"HRV_LFn": [0.25], "HRV_HFn": [0.0]})
    fd.attrs["method"] = "welch"
    rows = _pdf({"hrv_freq": fd})
    assert rows["LF n.u."] == "25.0 %"
    assert "HF n.u." not in rows         # as before: a 0 share is not listed


def test_pdf_rows_intervals_skip_unmeasured_columns():
    ivl = pd.DataFrame({"PR_ms": [np.nan] * 5, "QT_ms": [50.0, 52.0, 54.0, 56.0, 58.0]})
    rows = _pdf({"intervals": ivl})
    assert "PR" not in rows
    assert rows["QT"] == "54.0 ms"
    assert rows["QT IQR (beat-to-beat)"] == "4.0 ms"
