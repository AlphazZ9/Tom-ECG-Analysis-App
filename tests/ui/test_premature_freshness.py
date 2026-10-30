"""
Abnormal-event results in the exports must match the peaks they sit next
to: the classification records the windowed peaks and removed beats it ran
on, and the exports skip (and say so) a classification made on an earlier
peak set, artifact correction or analysis window. Also: very premature
beats removed as duplicate / non-physiological and legacy removals of
unknown type reach the summary through the stub app.
"""
from types import SimpleNamespace

import numpy as np
from openpyxl import Workbook

from ecg.core.detection import correct_rr_artifacts
from ecg.io.export import ExcelExporter, premature_beats_frame
from ecg.ui.export_controller import ExportController

from tests.helpers import rpeaks_from_rr
from tests.ui.test_premature_events_view import FS, _rr_mixture, _run, _stub_app


def _classified(peaks, report=None):
    app, statuses = _stub_app(peaks, artifact_report=report)
    ctrl, _ = _run(app)
    app.analysis_ctrl = ctrl
    return app, ctrl, statuses


def test_fresh_classification_is_current_and_exported():
    app, ctrl, _ = _classified(rpeaks_from_rr(_rr_mixture(), FS))
    assert ctrl.arrhythmia_is_current()
    events, summary, stale = ExportController(app).arrhythmia_export_data()
    assert not stale and events and summary["n_premature"] == 7
    s = app.analysis.premature_summary
    rp = app.detection.rpeaks_ok
    assert s["window_start_s"] == rp[0] / FS and s["window_end_s"] == rp[-1] / FS
    df = premature_beats_frame(s)
    assert df.loc[0, "Window_start_s"] == round(rp[0] / FS, 2)
    assert "Window_end_s" in df.columns


def test_peak_edit_makes_the_classification_stale():
    app, ctrl, _ = _classified(rpeaks_from_rr(_rr_mixture(), FS))
    app.detection.rpeaks_ok = np.delete(app.detection.rpeaks_ok, 50)
    assert not ctrl.arrhythmia_is_current()
    events, summary, stale = ExportController(app).arrhythmia_export_data()
    assert stale and events is None and summary is None


def test_window_change_makes_the_classification_stale():
    app, ctrl, _ = _classified(rpeaks_from_rr(_rr_mixture(), FS))
    full = app.detection.rpeaks_ok
    app._windowed_peaks = lambda: full[10:]
    assert not ctrl.arrhythmia_is_current()


def test_artifact_correction_change_makes_the_classification_stale():
    full = rpeaks_from_rr(_rr_mixture(), FS)
    corrected, report = correct_rr_artifacts(full, FS)
    app, ctrl, _ = _classified(corrected, report)
    assert ctrl.arrhythmia_is_current()
    app.analysis.artifact_report = None            # correction dropped
    assert not ctrl.arrhythmia_is_current()


def test_no_classification_is_never_stale():
    app, _ = _stub_app(rpeaks_from_rr(_rr_mixture(), FS))
    from ecg.ui.analysis_controller import AnalysisController
    app.analysis_ctrl = AnalysisController(app)
    assert ExportController(app).arrhythmia_export_data() == (None, None, False)


def test_stale_note_goes_to_the_cover_without_sheets():
    wb = Workbook()
    wb.active.title = "Cover"
    ExcelExporter.add_abnormal_events_sheets(
        wb, None, None, stale_note=ExportController.ARRHYTHMIA_STALE_TEXT)
    assert wb.sheetnames == ["Cover"]
    values = [c.value for row in wb["Cover"].iter_rows() for c in row]
    assert ExportController.ARRHYTHMIA_STALE_TEXT in values


def test_very_premature_beat_removed_as_duplicate_counts_in_the_tab():
    full = rpeaks_from_rr([120.0] * 30 + [60.0, 180.0] + [120.0] * 30, FS)
    corrected, report = correct_rr_artifacts(full, FS)
    assert "ectopic" not in report["removed_types"]
    app, _, statuses = _classified(corrected, report)
    s = app.analysis.premature_summary
    assert (s["n_isolated"], s["n_removed"]) == (1, 1)
    assert "still counted" in statuses[-1]


def test_legacy_report_of_unknown_type_is_disclosed_in_the_tab():
    full = rpeaks_from_rr(_rr_mixture(), FS)
    corrected, report = correct_rr_artifacts(full, FS)
    legacy = {k: v for k, v in report.items() if k not in ("removed_types", "break_samples")}
    app, _, statuses = _classified(corrected, legacy)
    s = app.analysis.premature_summary
    assert s["n_removed_unknown_type"] == len(report["removed_samples"])
    assert s["n_premature"] == 7
    assert "unknown type" in app.lbl_arr_premature.text
    assert "re-run artifact correction" in statuses[-1]
    assert "N_removed_unknown_type" in premature_beats_frame(s).columns
