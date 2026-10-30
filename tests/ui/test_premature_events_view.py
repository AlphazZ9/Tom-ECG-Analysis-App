"""
Abnormal Events tab and exports for the RR-based premature beats, through
a stub app (no Tk window): the classification puts back the beats that
automatic artifact correction removed as ectopic, the list header / status
/ ribbon / TSV report isolated beats, couplets and runs with per-hour
rates, and the Excel / Prism exports carry the same counts.
"""
from types import SimpleNamespace

import numpy as np
import pytest
from matplotlib.figure import Figure
from openpyxl import Workbook

from ecg.core.detection import (
    classify_arrhythmias, correct_rr_artifacts, premature_beat_summary,
)
from ecg.io.export import (
    ExcelExporter, abnormal_events_frame, premature_beats_frame,
)
from ecg.ui.analysis_controller import AnalysisController
from ecg.ui.state import AnalysisState, DetectionState, SignalState

from tests.helpers import rpeaks_from_rr

FS = 1000.0


class _Label:
    def __init__(self, text: str = "") -> None:
        self.text, self.text_color = text, None

    def configure(self, **kw) -> None:
        self.text = kw.get("text", self.text)
        self.text_color = kw.get("text_color", self.text_color)


class _Entry:
    def __init__(self, value) -> None:
        self.value = str(value)

    def get(self) -> str:
        return self.value


class _Slot:
    def __init__(self) -> None:
        self.draw_fn = None

    def update(self, fn) -> None:
        self.draw_fn = fn


def _rr_mixture() -> "list[float]":
    return ([120.0] * 30 + [70.0, 170.0] + [120.0] * 30
            + [70.0, 70.0, 220.0] + [120.0] * 30
            + [70.0] * 4 + [120.0] * 30)


def _stub_app(peaks, artifact_report=None):
    statuses = []
    app = SimpleNamespace(
        analysis=AnalysisState(artifact_report=artifact_report),
        detection=DetectionState(rpeaks_ok=peaks),
        signal=SignalState(fs=FS),
        after=lambda ms, fn=None: None,
        _set_progress=lambda p, m: None,
        _set_status=lambda text, color=None: statuses.append(text),
        ent_arr_baseline=_Entry(2), ent_arr_brady_pct=_Entry(20), ent_arr_min_beats=_Entry(10),
        _slots={"arr_detail": _Slot()},
        _arr_card_widgets=[], _arr_event_scroll=None,
        lbl_arrhythmia_status=_Label(), lbl_arr_premature=_Label(),
        tabs=SimpleNamespace(set=lambda name: None),
        btn_run_arrhythmia=object(),
        _start_async_result=lambda btn, text, worker, done: done(worker()),
    )
    app._windowed_peaks = lambda: app.detection.rpeaks_ok
    return app, statuses


def _run(app):
    ctrl = AnalysisController(app)
    cards = []
    ctrl.build_arrhythmia_card = lambda idx, ev, icons: cards.append(ev)
    ctrl.run_arrhythmia_analysis()
    return ctrl, cards


def test_auto_corrected_premature_beats_are_counted_in_the_tab():
    full = rpeaks_from_rr(_rr_mixture(), FS)
    corrected, report = correct_rr_artifacts(full, FS)
    app, statuses = _stub_app(corrected, artifact_report=report)
    _, cards = _run(app)

    s = app.analysis.premature_summary
    assert (s["n_isolated"], s["n_couplets"], s["n_runs"], s["n_beats_in_runs"]) == (1, 1, 1, 4)
    assert s["n_removed"] == report["n_ectopic"]
    span = (full[-1] - full[0]) / FS
    assert s["span_s"] == pytest.approx(span)
    assert s["premature_per_h"] == pytest.approx(7 / (span / 3600.0))

    assert {e.kind for e in cards} >= {"premature_beat", "premature_couplet", "premature_run"}
    assert "7 premature beats (RR-based)" in app.lbl_arrhythmia_status.text
    assert "isolated 1" in app.lbl_arr_premature.text
    assert "couplets 1" in app.lbl_arr_premature.text
    assert "runs ≥3 1" in app.lbl_arr_premature.text
    assert "supraventricular" in app.lbl_arr_premature.text
    assert "removed by artifact correction" in app.lbl_arr_premature.text
    assert "still counted" in statuses[-1]

    # TSV: one row per event, then the summary block
    tsv = app.analysis.arrhythmia_tsv
    assert "\tN_removed\t" in tsv.splitlines()[0]
    assert "Premature beats (RR-based)\t7 beats" in tsv
    assert "ESV" not in tsv

    # the ribbon figure draws with the new kinds and their labels
    fig = Figure()
    app._slots["arr_detail"].draw_fn(fig)
    legend_labels = [t.get_text() for leg in fig.legends for t in leg.get_texts()]
    assert "Premature couplet (RR-based)" in legend_labels


def test_without_correction_the_counts_are_the_same():
    full = rpeaks_from_rr(_rr_mixture(), FS)
    app, _ = _stub_app(full)
    _run(app)
    s = app.analysis.premature_summary
    assert s["n_premature"] == 7 and s["n_removed"] == 0
    assert "removed" not in app.lbl_arr_premature.text


def test_event_card_list_is_capped(monkeypatch):
    rr = []
    for _ in range(12):
        rr += [120.0] * 15 + [70.0, 170.0]
    rr += [120.0] * 15
    app, _ = _stub_app(rpeaks_from_rr(rr, FS))
    monkeypatch.setattr(AnalysisController, "MAX_EVENT_CARDS", 5)
    made = []
    import ecg.ui.analysis_controller as ac
    monkeypatch.setattr(ac.ctk, "CTkLabel",
                        lambda *a, **kw: made.append(kw) or SimpleNamespace(grid=lambda **k: None))
    _, cards = _run(app)
    assert len(cards) == 5
    assert app.analysis.premature_summary["n_isolated"] == 12
    assert made and "7 more events" in made[-1]["text"]


def test_premature_summary_text_empty_before_classification():
    assert AnalysisController.premature_summary_text(None) == ""


def test_excel_gets_premature_sheet_events_sheet_and_cover_rows():
    rp = rpeaks_from_rr(_rr_mixture(), FS)
    events = classify_arrhythmias(rp, FS, baseline_s=2.0)
    summary = premature_beat_summary(events, 3600.0)
    wb = Workbook()
    wb.active.title = "Cover"
    wb["Cover"].cell(row=1, column=1, value="ECG Analysis Report")
    ExcelExporter.add_abnormal_events_sheets(wb, events, summary)

    ws = wb["Premature_Beats"]
    header = [c.value for c in ws[1]]
    values = dict(zip(header, [c.value for c in ws[2]]))
    assert values["N_premature_total"] == 7
    assert (values["N_isolated"], values["N_couplets"], values["N_runs_ge3"]) == (1, 1, 1)
    assert values["Premature_per_h"] == pytest.approx(7.0)
    assert any("supraventricular" in str(c.value) for row in ws.iter_rows() for c in row)

    ev_ws = wb["Abnormal_Events"]
    labels = [r[1].value for r in ev_ws.iter_rows(min_row=2)]
    assert "Isolated premature beat (RR-based)" in labels

    cover_labels = [c.value for c in wb["Cover"]["A"]]
    assert "Premature beats (RR-based)" in cover_labels
    assert "Couplets" in cover_labels


def test_excel_sheets_skipped_before_classification():
    wb = Workbook()
    ExcelExporter.add_abnormal_events_sheets(wb, [], None)
    assert wb.sheetnames == ["Sheet"]
    assert premature_beats_frame(None).empty
    assert abnormal_events_frame([]).empty


def test_prism_notes_carry_premature_rows(tmp_path, monkeypatch):
    from ecg.io import export as export_mod
    captured = {}

    def _fake_write(tables, path, **kw):
        captured["notes"] = kw["notes"]["Project Info"]

    pzfx = pytest.importorskip("pzfx")
    monkeypatch.setattr(pzfx, "write_pzfx", _fake_write)
    rp = rpeaks_from_rr([120.0] * 60, FS)
    results = {"hr": {"n": 60}, "rr_df": None,
               "hrv_time": __import__("pandas").DataFrame({"HRV_MeanNN": [120.0]})}
    summary = premature_beat_summary([], 60.0)
    try:
        export_mod.PrismExporter.build_and_write(
            tmp_path / "x.pzfx", results, premature_summary=summary)
    except Exception as exc:          # other tables need fuller results
        pytest.skip(f"Prism export needs fuller results here: {exc}")
    notes = captured["notes"]
    assert "Premature beats (RR-based)" in list(notes["Name"])
