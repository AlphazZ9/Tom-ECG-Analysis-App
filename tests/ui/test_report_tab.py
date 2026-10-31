# -*- coding: utf-8 -*-
"""
The Report tab without a window: module lines and Include, what an export holds
for what Include says, the agreement of Report, screens and Metrics sheet, and
the session's former Summary tab.
"""
import io
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest

from ecg.core.metric_table import (
    REPORT_AUTONOMIC, REPORT_EVENTS, REPORT_INTERVALS, REPORT_MODULES, REPORT_QUALITY,
    REPORT_RHYTHM, REPORT_TIME_COURSE, build_metric_table, filter_results, filter_table,
    format_metric, metric_text, report_key_results, report_module_of, rhythm_cards,
)
from ecg.ui.export_controller import ExportController
from ecg.ui.freshness import NOT_COMPUTED, OUT_OF_DATE, UP_TO_DATE, combine_states
from ecg.ui.report_view import (
    autonomic_method_text, events_baseline_text, hrv_bands_text, include_choices, left_out_labels,
    provenance_rows, report_lines, report_states, report_text,
)
from ecg.ui.session_controller import _TAB_KEY_MAP
from ecg.ui.state import AnalysisState, DetectionState, SignalState, UIState
from ecg.ui.tab_keys import TAB_MODULES, TAB_ORDER, TAB_REPORT


def _results(n: int = 400) -> dict:
    rr = 120.0 + np.sin(np.arange(n)) * 3.0
    t = np.cumsum(rr) / 1000.0
    fd = pd.DataFrame({"HRV_VLF": [1.5], "HRV_LF": [2.25], "HRV_HF": [3.5], "HRV_LFn": [0.4],
                       "HRV_HFn": [0.6], "HRV_LFHF": [0.6429]})
    fd.attrs.update(method="welch", bands=[(0.1, 0.4), (0.4, 1.5), (1.5, 5.0)])
    td = pd.DataFrame({"HRV_MeanNN": [120.0], "HRV_SDNN": [3.14159], "HRV_RMSSD": [2.5],
                       "HRV_pNN6": [10.0], "HRV_pNN20": [0.0]})
    nl = pd.DataFrame({"HRV_SD1": [1.8], "HRV_SD2": [3.9], "HRV_SampEn": [1.2]})
    nl.attrs["n_rr_valid"] = n
    ivl = pd.DataFrame({"PR_ms": [35.0, 36.0, 34.0, np.nan], "QRS_ms": [10.0, 11.0, 10.0, 9.0],
                        "QT_ms": [50.0, 52.0, 51.0, 49.0], "QTc_ms": [45.0, 47.0, 46.0, 44.0],
                        "RR_ms": [120.0] * 4, "R_peak_s": [1.0, 2.0, 3.0, 4.0]})
    return {"hr": {"n": n + 1, "n_valid": n, "n_excluded": 0, "mean": 500.0, "min": 480.0,
                   "max": 520.0, "std": 2.0, "rr_min_ms": 66.7, "rr_max_ms": 333.3},
            "rr_ms": rr, "rr_valid": np.ones(n, dtype=bool), "rr_t_s": t,
            "rr_df": pd.DataFrame({"Time_s": t, "RR_ms": rr, "HR_bpm": 60000.0 / rr}),
            "hrv_time": td, "hrv_freq": fd, "hrv_nonlin": nl, "intervals": ivl,
            "beat_corr": np.full(n, 0.97)}


def _epochs() -> pd.DataFrame:
    return pd.DataFrame({"Epoch_start_s": [0.0, 30.0], "Epoch_end_s": [30.0, 60.0],
                         "HR_mean": [500.0, 501.0], "SDNN": [3.0, 3.1], "N_beats": [250, 250],
                         "N_valid": [249, 249], "N_excluded": [0, 0]})


def _rolling() -> pd.DataFrame:
    return pd.DataFrame({"t_mid": [15.0, 20.0, 25.0], "HR": [500.0, 500.0, 501.0],
                         "SDNN": [3.0, 3.0, 3.1], "n_beats": [250, 250, 250]})


class _Fresh:
    def __init__(self, states):
        self._states = states

    def states(self):
        return dict(self._states)

    def status(self, module):
        return SimpleNamespace(state=self._states[module],
                               reasons=("peaks changed",) if self._states[module] == OUT_OF_DATE else ())


def _all_states(**over):
    states = {m: NOT_COMPUTED for m in ("autonomic", "event_locked", "events")}
    states.update({m: UP_TO_DATE for m in ("core", "frequency", "nonlinear", "intervals",
                                           "epochs", "rolling")})
    states.update(over)
    return states


def _app(states=None, include=None, results=None):
    a = AnalysisState(results=results if results is not None else _results(),
                      analysed_window=(0.0, 0.0))
    a.epoch_df = _epochs()
    a.rolling_hrv_df = _rolling()
    t = np.arange(0, 60, 0.01)
    sig = np.sin(t)
    app = SimpleNamespace(
        analysis=a, detection=DetectionState(rpeaks_ok=np.arange(10, 5000, 50), sig_quality=88),
        signal=SignalState(fs=100, filtered=sig, raw_norm=sig, time=t, filepath="C:/data/m1.mat"),
        ui=UIState(report_include=include or {}),
        ent_subject=SimpleNamespace(get=lambda: "m1"), _qtc_formula=lambda: "mitchell",
        session=SimpleNamespace(recording_notes=""),
    )
    if states is not None:
        app.freshness = _Fresh(states)
    return app


def _sheets(app):
    buf = io.BytesIO()
    ExportController(app).write_excel(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


# ── tabs and the former Summary ──────────────────────────────────────────

def test_report_replaced_summary_in_the_tab_bar():
    assert TAB_ORDER == ("Signal & Detection", "Rhythm & HRV", "Time course",
                         "Morphology & Intervals", "Events", "Report")
    assert TAB_MODULES[TAB_REPORT] == ("core",)


@pytest.mark.parametrize("saved", ["Summary", "📋 Summary"])
def test_an_old_session_saved_on_summary_opens_on_report(saved):
    assert _TAB_KEY_MAP[saved] == TAB_REPORT


# ── state and Include ────────────────────────────────────────────────────

def test_lines_take_the_state_of_their_modules():
    states = report_states(_all_states(rolling=OUT_OF_DATE))
    assert states[REPORT_RHYTHM] == UP_TO_DATE
    assert states[REPORT_TIME_COURSE] == OUT_OF_DATE
    assert states[REPORT_EVENTS] == NOT_COMPUTED
    assert [m.label for m in REPORT_MODULES] == [
        "Signal quality", "Rhythm & HRV", "Time course", "Autonomic CSI/CPI", "Intervals", "Events"]


def test_include_is_on_by_default_only_when_up_to_date():
    got = include_choices({"a": UP_TO_DATE, "b": OUT_OF_DATE, "c": NOT_COMPUTED})
    assert got == {"a": True, "b": False, "c": False}


def test_a_tick_holds_until_the_state_changes():
    # Out-of-date values go out only as long as they were asked for in that state.
    ticks = {"b": [True, OUT_OF_DATE], "a": [False, UP_TO_DATE], "c": [True, UP_TO_DATE]}
    got = include_choices({"a": UP_TO_DATE, "b": OUT_OF_DATE, "c": NOT_COMPUTED}, ticks)
    assert got == {"a": False, "b": True, "c": False}
    assert include_choices({"b": UP_TO_DATE}, ticks)["b"] is True
    # A tick from when it was up to date does not send it once out of date.
    assert include_choices({"a": OUT_OF_DATE}, {"a": [True, UP_TO_DATE]})["a"] is False


def test_left_out_lists_only_what_was_computed():
    per_line = {REPORT_TIME_COURSE: OUT_OF_DATE, REPORT_EVENTS: NOT_COMPUTED, REPORT_RHYTHM: UP_TO_DATE}
    got = left_out_labels(per_line, {REPORT_TIME_COURSE: False, REPORT_EVENTS: False, REPORT_RHYTHM: True})
    assert got == ["Time course (out of date)"]


# ── key results and provenance ───────────────────────────────────────────

def test_every_value_on_report_matches_its_screen_and_its_metrics_row():
    r = _results()
    t = build_metric_table(r, sig_quality=88, epoch_df=_epochs(), rolling_df=_rolling())
    keys = report_key_results(t)
    cards = {c.key: c for c in rhythm_cards(t)}
    frame = t.to_frame().set_index("Key")

    assert keys[REPORT_RHYTHM] == (f"HR {cards['hr'].value} bpm · SDNN {cards['sdnn'].value} ms · "
                                   f"LF/HF {cards['lf_hf'].value}")
    for text_key, card in (("hr.mean", "hr"), ("hrv_time.HRV_SDNN", "sdnn"), ("hrv_freq.HRV_LFHF", "lf_hf")):
        # The Metrics sheet keeps the raw value; the screens round it the same way.
        assert format_metric(text_key, frame.loc[text_key, "Value"], t.row(text_key)) == cards[card].value
    assert keys[REPORT_QUALITY] == (f"score {frame.loc['quality.score', 'Value']:.0f} · "
                                    f"corr {frame.loc['quality.mean_corr', 'Value']:.3f}")
    assert keys[REPORT_TIME_COURSE] == (f"{frame.loc['epochs.status', 'Basis_n']} epochs · "
                                        f"{frame.loc['rolling.status', 'Basis_n']} sliding windows")
    for col, name in (("PR_ms", "PR"), ("QRS_ms", "QRS"), ("QT_ms", "QT"), ("QTc_ms", "QTc")):
        assert f"{name} {frame.loc[f'intervals.{col}', 'Value']:.1f}" in keys[REPORT_INTERVALS]
    measured = frame.loc["intervals.beats_measured"]
    assert f"{measured['Value']}/{measured['Basis_total']} beats measured" in keys[REPORT_INTERVALS]
    assert metric_text(t, "hr.mean") == "500"


def test_a_line_without_a_value_is_a_dash_and_no_number_is_made_up():
    keys = report_key_results(build_metric_table(None))
    assert set(keys.values()) == {"—"}


def test_an_out_of_date_line_says_why_next_to_its_value():
    t = build_metric_table(_results(), epoch_df=_epochs())
    per_line = report_states(_all_states(epochs=OUT_OF_DATE))
    lines = {x.key: x for x in report_lines(t, per_line, include_choices(per_line),
                                            {REPORT_TIME_COURSE: ("peaks changed",)})}
    assert lines[REPORT_TIME_COURSE].state == OUT_OF_DATE
    assert lines[REPORT_TIME_COURSE].result == "2 epochs · peaks changed"
    assert not lines[REPORT_TIME_COURSE].include and lines[REPORT_TIME_COURSE].can_include
    assert not lines[REPORT_EVENTS].can_include


def test_provenance_reads_what_the_results_were_computed_on():
    r = _results()
    t = build_metric_table(r, artifact_report=None)
    rows = dict(provenance_rows(t, r, "C:/data/m1.mat", "m1", (0.0, 0.0), 60.0, "none",
                                "mitchell", "Robust · 3 s · ks 3 / kp 12 (rat)"))
    assert rows["File"] == "m1.mat" and rows["Subject"] == "m1"
    assert rows["Analysis window"] == "0–60 s · 401 beats"
    assert rows["RR bounds"] == "66.7–333.3 ms"
    assert rows["Excluded RR"] == "0 RR excluded"
    assert rows["HRV bands"] == "LF 0.4–1.5 · HF 1.5–5 Hz"
    assert rows["QTc formula"].startswith("Mitchell")
    assert rows["Autonomic"].endswith("(rat)")


def test_method_and_baseline_texts():
    assert autonomic_method_text(None) == ""
    assert autonomic_method_text({"method_label": "Robust", "window_s": 3.0, "ks": 3.0, "kp": 12.0}
                                 ) == "Robust · 3 s · ks 3 / kp 12 (rat)"
    assert events_baseline_text({"baseline": {"baseline_hr": 509.4}}) == "baseline 509 bpm"
    assert events_baseline_text({"baseline": {"baseline_hr": float("nan")}}) == ""
    assert hrv_bands_text({}) == "—"


# ── what goes out ────────────────────────────────────────────────────────

def test_filter_results_is_the_same_dict_with_everything_in():
    r = _results()
    assert filter_results(r, {m.key: True for m in REPORT_MODULES}) is r
    no_rhythm = filter_results(r, {REPORT_RHYTHM: False})
    assert "hrv_time" not in no_rhythm and "rr_ms" not in no_rhythm and "hr" in no_rhythm
    assert "intervals" in no_rhythm and "intervals" not in filter_results(r, {REPORT_INTERVALS: False})
    assert r["hrv_time"] is not None      # the results themselves are untouched


def test_every_table_row_belongs_to_a_line_or_to_the_provenance():
    t = build_metric_table(_results(), sig_quality=88, epoch_df=_epochs(), rolling_df=_rolling(),
                           artifact_report={"n_in": 3, "n_out": 2})
    keys = {m.key for m in REPORT_MODULES}
    for row in t.rows:
        line = report_module_of(row)
        assert line in keys or row.key.startswith("artifact.") or row.key == "core.duration_s"
    off = filter_table(t, {REPORT_QUALITY: False, REPORT_INTERVALS: False})
    assert not off.has("quality.score") and not off.has("intervals.QT_ms")
    assert off.has("artifact.removed") and off.has("hr.mean")


def test_an_app_without_the_freshness_service_exports_everything():
    wb = _sheets(_app(states=None))
    assert {"Cover", "HR_Statistics", "RR_Timeseries", "HRV_Time", "HRV_Frequency", "HRV_NonLinear",
            "ECG_Intervals", "Epochs", "Metrics"} <= set(wb.sheetnames)


def test_all_included_has_every_sheet_and_no_left_out_row():
    wb = _sheets(_app(_all_states()))
    assert {"HR_Statistics", "RR_Timeseries", "HRV_Time", "ECG_Intervals", "Epochs", "Metrics"} <= set(wb.sheetnames)
    labels = [c.value for c in wb["Cover"]["A"]]
    assert "Left out of this export" not in labels and "HR mean" in labels


def test_a_stale_module_stays_out_by_default_and_the_cover_says_so():
    wb = _sheets(_app(_all_states(epochs=OUT_OF_DATE)))
    assert "Epochs" not in wb.sheetnames
    cover = {r[0].value: r[1].value for r in wb["Cover"].iter_rows() if len(r) > 1}
    assert cover["Left out of this export"] == "Time course (out of date)"
    assert not [k for k in (c.value for c in wb["Metrics"]["A"]) if str(k).startswith("epochs.")]


def test_ticking_a_stale_module_sends_it():
    states = _all_states(epochs=OUT_OF_DATE)
    wb = _sheets(_app(states, include={REPORT_TIME_COURSE: [True, OUT_OF_DATE]}))
    assert "Epochs" in wb.sheetnames


def test_unticking_rhythm_and_quality_takes_their_sheets_and_cover_rows_out():
    include = {REPORT_RHYTHM: [False, UP_TO_DATE], REPORT_QUALITY: [False, UP_TO_DATE]}
    wb = _sheets(_app(_all_states(), include=include))
    for sheet in ("HR_Statistics", "RR_Timeseries", "HRV_Time", "HRV_Frequency", "HRV_NonLinear"):
        assert sheet not in wb.sheetnames
    assert "ECG_Signal" in wb.sheetnames and "ECG_Intervals" in wb.sheetnames
    labels = [c.value for c in wb["Cover"]["A"]]
    for label in ("HR mean", "HR range", "Signal quality", "Frequency HRV"):
        assert label not in labels
    assert "N beats" in labels


def test_unticking_intervals_takes_the_sheet_and_the_definitions_out():
    wb = _sheets(_app(_all_states(), include={REPORT_INTERVALS: [False, UP_TO_DATE]}))
    assert "ECG_Intervals" not in wb.sheetnames
    assert "Interval definitions" not in [c.value for c in wb["Cover"]["A"]]


def test_a_module_never_computed_is_not_listed_as_left_out():
    results = _results()
    results["intervals"] = pd.DataFrame({"RR_ms": [120.0]})   # analyse_core's placeholder
    app = _app(_all_states(intervals=NOT_COMPUTED), results=results)
    data = ExportController(app).export_data()
    assert data.left_out == ""
    assert "intervals" in data.results


def test_prism_gets_what_include_lets_out(monkeypatch, tmp_path):
    from ecg.io import export as export_module
    captured = {}
    monkeypatch.setattr(export_module.PrismExporter, "build_and_write",
                        classmethod(lambda cls, **kw: captured.update(kw) or 1))
    import ecg.ui.export_controller as ctrl
    monkeypatch.setattr(ctrl.filedialog, "asksaveasfilename", lambda **k: str(tmp_path / "x.pzfx"))
    app = _app(_all_states(rolling=OUT_OF_DATE), include={REPORT_INTERVALS: [False, UP_TO_DATE]})
    app._set_status = lambda *a, **k: None
    ExportController(app).export_prism()
    assert captured["rolling_hrv_df"] is None and captured["epoch_df"] is None
    assert "intervals" not in captured["results"]
    assert captured["left_out"] == "Time course (out of date), Intervals (up to date)"


# ── text report ──────────────────────────────────────────────────────────

def _text(app):
    return ExportController(app).report_text()


def test_text_report_holds_module_states_and_no_invented_number():
    text = _text(_app(_all_states(rolling=OUT_OF_DATE)))
    assert "ECG ANALYSIS  —  Report" in text and "Summary" not in text
    assert "WHAT IS IN THIS REPORT" in text
    assert "(left out)" in text and "HEART RATE" in text


def test_text_report_follows_include():
    include = {REPORT_RHYTHM: [False, UP_TO_DATE], REPORT_INTERVALS: [False, UP_TO_DATE]}
    text = _text(_app(_all_states(), include=include))
    for head in ("HEART RATE", "HRV — TIME DOMAIN", "HRV — NON-LINEAR", "ECG INTERVALS", "RR ASYMMETRY"):
        assert head not in text
    assert "PROVENANCE" in text


def test_text_numbers_do_not_depend_on_the_state_block():
    """The block added for the Report holds no number of its own."""
    r = _results()
    lines = report_lines(build_metric_table(None), report_states(_all_states()),
                         {m.key: True for m in REPORT_MODULES})
    base = report_text(r, when=pd.Timestamp("2026-01-02 03:04"))
    with_block = report_text(r, when=pd.Timestamp("2026-01-02 03:04"), lines=lines)
    head, block = with_block.split("  WHAT IS IN THIS REPORT\n")
    assert base.startswith(head.rstrip("\n")[:40])
    assert not any(ch.isdigit() for ch in block.split("\n\n")[0])


def test_combine_states_of_a_report_line():
    assert combine_states([NOT_COMPUTED, UP_TO_DATE]) == UP_TO_DATE
    assert combine_states([NOT_COMPUTED, NOT_COMPUTED]) == NOT_COMPUTED
    assert combine_states([OUT_OF_DATE, UP_TO_DATE]) == OUT_OF_DATE
    assert REPORT_AUTONOMIC in {m.key for m in REPORT_MODULES}
