"""
UI <-> computation links, end to end through the controllers with a stub
app (no Tk window), on a synthetic ECG whose first peak is too close to the
signal start for a full beat window (the edge-beat case that broke the
AV-delay check and shifted the quality curve):

- Intervals -> AV delay: analyse_intervals rows reach
  AnalysisController.run_arrhythmia_analysis matched to the peaks;
- Summary: PlotController.plot_summary writes the report, the metric tiles,
  the RR-asymmetry block and a quality-over-time curve timed by the beats
  actually scored;
- Excel: ExportController.write_excel carries the ECG_Intervals and
  Abnormal_Events sheets of the same results.
"""
import io
import re
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook

from ecg.core import analysis
from ecg.core.analysis import analyse_core, analyse_intervals
from ecg.core.wave_template import WaveTemplate
from ecg.ui.analysis_controller import AnalysisController
from ecg.ui.export_controller import ExportController
from ecg.ui.plot_controller import PlotController, _df_val
from ecg.ui.state import AnalysisState, DetectionState, SignalState

from tests.helpers import make_synth_ecg

pytestmark = pytest.mark.skipif(not analysis.NK_AVAILABLE, reason="neurokit2 not installed")


@pytest.fixture(autouse=True)
def _no_template_writes(monkeypatch):
    def _refuse(self):
        raise AssertionError("WaveTemplate.save() must not run in tests")
    monkeypatch.setattr(WaveTemplate, "save", _refuse)


@pytest.fixture(autouse=True)
def _no_tk_widgets(monkeypatch):
    """The "No abnormal events" placeholder is a CTkLabel: no Tk here."""
    from ecg.ui import analysis_controller as ac
    monkeypatch.setattr(ac.ctk, "CTkLabel",
                        lambda *a, **kw: SimpleNamespace(grid=lambda **k: None,
                                                         destroy=lambda: None))


class _Label:
    def __init__(self, text: str = "") -> None:
        self.text, self.text_color = text, None

    def configure(self, **kw) -> None:
        self.text = kw.get("text", self.text)
        self.text_color = kw.get("text_color", self.text_color)

    def cget(self, key):
        return self.text


class _Slot:
    def __init__(self) -> None:
        self.draw_fn = None

    def update(self, fn) -> None:
        self.draw_fn = fn


class _Entry:
    def __init__(self, value) -> None:
        self.value = str(value)

    def get(self) -> str:
        return self.value


@pytest.fixture(scope="module")
def recording():
    sig, fs, rp = make_synth_ecg(n_beats=60, seed=2)
    crop = rp[0] - 10                    # first peak 5 ms after the start
    sig, rp = sig[crop:], rp - crop
    r = analyse_core(sig, rp, fs)
    wt = WaveTemplate()
    wt.confirmed = True
    ivl = analyse_intervals(sig, rp, fs, r["rr_ms"], wave_template=wt)
    # The synthetic beats have no P wave: keep analyse_intervals' rows (one
    # per peak with a full beat window, timed by R_peak_s) with a normal PR.
    ivl["PR_ms"] = 40.0
    r["intervals"] = ivl
    return sig, fs, rp, r


def _stub_app(sig, fs, rp, r):
    statuses: list = []
    texts: dict = {}
    slots = {k: _Slot() for k in ("arr_detail", "sum_rr", "sum_poincare",
                                  "sum_asymmetry", "sum_quality_time")}
    app = SimpleNamespace(
        analysis=AnalysisState(results=r, analysed_window=(0.0, 0.0)),
        detection=DetectionState(rpeaks_ok=rp, sig_quality=90),
        signal=SignalState(fs=fs, filtered=sig, time=np.arange(len(sig)) / fs,
                           filepath="C:/data/m1.mat"),
        after=lambda ms, fn=None: None,
        _set_progress=lambda p, m: None,
        _set_status=lambda text, color=None: statuses.append(text),
        ent_arr_baseline=_Entry(2), ent_arr_brady_pct=_Entry(20), ent_arr_min_beats=_Entry(10),
        ent_subject=_Entry("m1"),
        _slots=slots, _arr_card_widgets=[], _arr_event_scroll=None,
        lbl_arrhythmia_status=_Label(), lbl_arr_premature=_Label(), lbl_sum_verdict=None,
        tabs=SimpleNamespace(set=lambda name: None),
        btn_run_arrhythmia=object(),
        _start_async_result=lambda btn, text, worker, done: done(worker()),
        _safe_df_val=_df_val,
        _sum_metric_vals={k: _Label() for k in ("hr_mean", "pr", "qrs", "qt", "qtc", "pnn6")},
        _sum_quality_vals={k: _Label() for k in ("sq_noisy_time", "sq_badbeats")},
        txt_sum="txt_sum",
        _set_textbox=lambda widget, text, tsv=None: texts.__setitem__(widget, text),
        _qtc_formula=lambda: "mitchell",
    )
    app._windowed_peaks = lambda: app.detection.rpeaks_ok
    return app, statuses, texts


def _run_arrhythmia(app):
    ctrl = AnalysisController(app)
    cards: list = []
    ctrl.build_arrhythmia_card = lambda idx, ev, icons: cards.append(ev)
    ctrl.run_arrhythmia_analysis()
    app.analysis_ctrl = ctrl
    return app.analysis.arrhythmia_events


# ── Intervals -> AV delay ───────────────────────────────────────────────────

def test_interval_rows_drive_the_av_delay_check(recording):
    sig, fs, rp, r = recording
    ivl = r["intervals"]
    assert len(ivl) < len(rp)                          # edge beat has no row
    assert np.isclose(ivl["R_peak_s"].iloc[0], rp[1] / fs)    # rows start at peak 2

    app, statuses, _ = _stub_app(sig, fs, rp, dict(r))
    events = _run_arrhythmia(app)
    assert "AV-delay check on" in statuses[-1]
    assert not any(e.kind == "av_delay" for e in events)

    # A sustained PR prolongation on rows 20..39 (peaks 21..40: the rows
    # start at the second peak) becomes one AV-delay event on those beats.
    long_pr = ivl.copy()
    long_pr.loc[long_pr.index[20:40], "PR_ms"] = 70.0
    app, statuses, _ = _stub_app(sig, fs, rp, dict(r, intervals=long_pr))
    av = [e for e in _run_arrhythmia(app) if e.kind == "av_delay"]
    assert len(av) == 1
    assert av[0].n_beats == 20
    row_peak = int(np.round(long_pr["R_peak_s"].iloc[20] * fs))
    assert av[0].t_start == pytest.approx(row_peak / fs)


def test_intervals_from_another_window_skip_the_av_check(recording):
    sig, fs, rp, r = recording
    app, statuses, _ = _stub_app(sig, fs, rp, dict(r))
    app._windowed_peaks = lambda: rp + 40          # no row matches these peaks
    _run_arrhythmia(app)
    assert "AV-delay check skipped — run Intervals on the current window" in statuses[-1]

    sparse = r["intervals"].copy()
    sparse.loc[sparse.index[5:], "PR_ms"] = np.nan
    app, statuses, _ = _stub_app(sig, fs, rp, dict(r, intervals=sparse))
    _run_arrhythmia(app)
    assert "PR measured on < 50% of beats" in statuses[-1]


# ── Summary ────────────────────────────────────────────────────────────────

def test_summary_generation(recording):
    sig, fs, rp, r = recording
    app, _, texts = _stub_app(sig, fs, rp, dict(r))
    PlotController(app).plot_summary(app.analysis.results)

    report = texts["txt_sum"]
    ivl = r["intervals"]
    for col, label in (("PR_ms", "PR"), ("QRS_ms", "QRS"), ("QT_ms", "QT")):
        med = pd.to_numeric(ivl[col], errors="coerce").dropna().median()
        line = next(ln for ln in report.splitlines()
                    if re.match(rf"\s*{label}\s+\d", ln))
        assert f"{med:.1f} ±" in line
    assert "Porta index (PI)" in report and "Guzik index (GI)" in report
    assert app._sum_metric_vals["pr"].text not in ("", "—")
    assert app._sum_metric_vals["qtc"].text not in ("", "—")
    assert app._sum_quality_vals["sq_noisy_time"].text.endswith("%")

    # Quality-over-time scatter: one point per scored beat, at its time.
    fig = Figure()
    app._slots["sum_quality_time"].draw_fn(fig)
    offsets = fig.axes[0].collections[0].get_offsets()
    np.testing.assert_allclose(np.asarray(offsets)[:, 0], r["beat_corr_t_s"])
    np.testing.assert_allclose(np.asarray(offsets)[:, 0], rp[1:] / fs)


# ── Excel export of intervals / events ─────────────────────────────────────

def test_excel_export_carries_intervals_and_events(recording):
    sig, fs, rp, r = recording
    long_pr = r["intervals"].copy()
    long_pr.loc[long_pr.index[20:40], "PR_ms"] = 70.0
    app, _, _ = _stub_app(sig, fs, rp, dict(r, intervals=long_pr))
    events = _run_arrhythmia(app)
    assert any(e.kind == "av_delay" for e in events)

    buf = io.BytesIO()
    ExportController(app).write_excel(buf)
    buf.seek(0)
    wb = load_workbook(buf)

    ws = wb["ECG_Intervals"]
    header = [c.value for c in ws[1]]
    for col in ("PR_ms", "QRS_ms", "QT_ms", "QTc_ms"):
        assert col in header
    pr_col = header.index("PR_ms") + 1
    pr_xl = [ws.cell(row=i, column=pr_col).value for i in range(2, len(long_pr) + 2)]
    np.testing.assert_allclose(np.array(pr_xl, dtype=float),
                               long_pr["PR_ms"].to_numpy(dtype=float), equal_nan=True)

    ev_ws = wb["Abnormal_Events"]
    assert ev_ws.max_row - 1 == len(events)
    ev_text = " ".join(str(c.value) for row in ev_ws.iter_rows() for c in row)
    assert "AV delay" in ev_text or "av_delay" in ev_text

    cover = {wb["Cover"].cell(row=i, column=1).value: wb["Cover"].cell(row=i, column=2).value
             for i in range(1, wb["Cover"].max_row + 1)}
    assert any(str(k).startswith("Interval definitions") for k in cover)
