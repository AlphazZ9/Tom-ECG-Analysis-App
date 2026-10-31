# -*- coding: utf-8 -*-
"""
ecg.ui.compare_segments
-----------------------
ECGApp Compare Segments dialog.
Mixin of ECGApp (see ecg/ui/app.py).
"""
from __future__ import annotations

import logging
import os
from datetime import datetime
from typing import Any, TYPE_CHECKING
import customtkinter as ctk  # type: ignore[import-untyped]
from tkinter import filedialog, messagebox
from matplotlib.figure import Figure
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from ecg.core.models import MouseECG
from ecg.ui.theme import (
    BG, PANEL, CARD, BORDER, BORDER2, TEXT, MUTED, LIGHT, PLOT, RED, BLUE,
    GREEN, ORANGE, TEAL, TEAL_DARK, GREEN_DARK, ORANGE_DARK, SPACE_XS, SPACE_S,
    SPACE_M, SPACE_L, FONT_LABEL, FONT_SMALL, FONT_KPI_VALUE, FONT_KPI_LABEL,
    FONT_BTN_PRIMARY, FONT_BTN_SEC, FONT_SIDEBAR_HDR, FONT_MICRO, FONT_BADGE,
    FONT_SUBSECTION,
)
from ecg.ui.plots import CanvasSlot, style_axes

# Mixin of ECGApp: `self: ECGApp` shows pyright the whole app surface, so a
# typo in self.xxx is reported. It is deliberately not a supertype of this
# class, which is the one check switched off here.
# pyright: reportGeneralTypeIssues=false
if TYPE_CHECKING:
    from ecg.ui.app import ECGApp

log = logging.getLogger("ecg")


class CompareSegmentsMixin:
    """ECGApp Compare Segments dialog."""

    # ════════════════════════════════════════════════════════
    #  COMPARE SEGMENTS
    # ════════════════════════════════════════════════════════

    # Rows of the Compare Segments table and export: (key in the
    # compute_segment_stats / compare_segments dict, label, unit). Mean HR
    # is 60000 / mean NN; the block rows are the non-pseudo-replicated
    # summary the descriptive comparison is based on.
    _SEGMENT_TABLE_METRICS: "list[tuple[str, str, str]]" = [
        ("n_beats",       "Beats",                      ""),
        ("duration_s",    "Duration",                   "s"),
        ("n_blocks",      "Blocks used (n)",            ""),
        ("hr_block_mean", "HR, mean of block means",    "bpm"),
        ("hr_block_sd",   "HR, SD of block means",      "bpm"),
        ("hr_mean",       "Mean HR (60000/mean NN)",    "bpm"),
        ("hr_sd",         "Instantaneous HR SD",        "bpm"),
        ("hr_min",        "Min inst. HR (2nd %ile)",    "bpm"),
        ("hr_max",        "Max inst. HR (98th %ile)",   "bpm"),
        ("rr_mean",       "Mean RR",                    "ms"),
        ("rr_sd",         "RR SD",                      "ms"),
        ("rr_cv",         "RR CV",                      "%"),
        ("sdnn",          "SDNN",                       "ms"),
        ("rmssd",         "RMSSD",                      "ms"),
        ("pnn6",          "pNN6",                       "%"),
        ("lf",            "LF power",                   "ms²"),
        ("hf",            "HF power",                   "ms²"),
        ("lf_nu",         "LF (n.u.)",                  "%"),
        ("hf_nu",         "HF (n.u.)",                  "%"),
        ("lf_hf",         "LF/HF",                      ""),
        ("sd1",           "Poincaré SD1",               "ms"),
        ("sd2",           "Poincaré SD2",               "ms"),
        ("sd_ratio",      "SD1/SD2",                    ""),
    ]

    def _open_compare_segments(self: ECGApp) -> None:
        """Open the Segment Comparison window (Time course > Event-locked, Ctrl+M).

        Two configurable time windows (A / B) → side-by-side metric table
        + superimposed RR tachogram.  All stats run on a background thread.
        """
        if self._signal_flt is None or self._rpeaks_ok is None or self._fs is None:
            messagebox.showwarning("No data", "Open a file and click Detect Peaks first.")
            return

        sig    = self._signal_flt
        rp     = self._rpeaks_ok.copy()
        fs     = self._fs
        dur    = float(len(sig)) / fs
        vlf, lf, hf = self._get_freq_bands()
        sub    = self.ent_subject.get().strip() if self.ent_subject else "subject"

        win = ctk.CTkToplevel(self)
        win.title("Compare Segments")
        win.geometry("1060x720")
        win.configure(fg_color=PANEL)
        win.resizable(True, True)

        hdr = ctk.CTkFrame(win, fg_color=PANEL, corner_radius=0)
        hdr.pack(fill="x")
        ctk.CTkLabel(hdr, text="⚖  Segment Comparison",
                     font=FONT_KPI_VALUE, text_color=TEXT,
                     anchor="w").pack(side="left", padx=SPACE_L, pady=SPACE_M)
        ctk.CTkLabel(hdr,
                     text=f"{sub}  ·  {dur:.1f} s  ·  {len(rp)} peaks",
                     font=FONT_SMALL, text_color=MUTED).pack(side="left", padx=(0, SPACE_L))
        ctk.CTkButton(hdr, text="✗  Close", command=win.destroy,
                      fg_color=BORDER, hover_color=BORDER2, text_color=MUTED,
                      font=FONT_BTN_SEC, height=28, corner_radius=6).pack(
            side="right", padx=SPACE_L, pady=SPACE_M)

        # ── Quick fill: before / after an event ─────────────────────────
        # Packed above the two segment cards now, populated with widgets
        # further down once seg_entries exists -- pack() position is set
        # by call order, so this still renders at the top of the window.
        quick_frame = ctk.CTkFrame(win, fg_color=CARD, corner_radius=8)
        quick_frame.pack(fill="x", padx=SPACE_L, pady=(SPACE_M, 0))

        seg_frame = ctk.CTkFrame(win, fg_color="transparent")
        seg_frame.pack(fill="x", padx=SPACE_L, pady=(SPACE_S, SPACE_S))

        seg_entries: "list[dict]" = []
        _colors = [BLUE, ORANGE_DARK]
        _names  = ["Segment A", "Segment B"]

        for i, (name, col) in enumerate(zip(_names, _colors)):
            card = ctk.CTkFrame(seg_frame, fg_color=CARD, corner_radius=8,
                                border_width=2, border_color=col)
            card.pack(side="left", fill="x", expand=True,
                      padx=(0, 8 if i == 0 else 0))
            ctk.CTkLabel(card, text=name, font=FONT_SIDEBAR_HDR,
                         text_color=col, anchor="w").pack(padx=SPACE_L, pady=(SPACE_M, SPACE_S), fill="x")
            row = ctk.CTkFrame(card, fg_color="transparent")
            row.pack(fill="x", padx=SPACE_L, pady=(0, SPACE_M))
            ctk.CTkLabel(row, text="From (s):", font=FONT_SMALL,
                         text_color=MUTED, width=62).pack(side="left")
            lo_e = ctk.CTkEntry(row, width=80, height=28, font=FONT_LABEL,
                                fg_color=BG, border_color=col, text_color=TEXT)
            lo_e.insert(0, str(int(dur * i / 2)))
            lo_e.pack(side="left", padx=(0, SPACE_M))
            ctk.CTkLabel(row, text="To (s):", font=FONT_SMALL,
                         text_color=MUTED, width=46).pack(side="left")
            hi_e = ctk.CTkEntry(row, width=80, height=28, font=FONT_LABEL,
                                fg_color=BG, border_color=col, text_color=TEXT)
            hi_e.insert(0, str(int(dur * (i + 1) / 2)))
            hi_e.pack(side="left", padx=(0, SPACE_M))
            lbl_e = ctk.CTkEntry(row, width=120, height=28, font=FONT_LABEL,
                                 fg_color=BG, border_color=BORDER2, text_color=TEXT,
                                 placeholder_text="Label (optional)")
            lbl_e.insert(0, name)
            lbl_e.pack(side="left")
            seg_entries.append({"lo": lo_e, "hi": hi_e, "lbl": lbl_e, "color": col})

        # ── Quick-fill content (frame already packed above the cards) ────
        ctk.CTkLabel(quick_frame, text="⚡  Quick fill — before / after an event",
                     font=FONT_SUBSECTION, text_color=TEAL, anchor="w"
                     ).pack(padx=SPACE_L, pady=(SPACE_S, 2), anchor="w")
        qrow = ctk.CTkFrame(quick_frame, fg_color="transparent")
        qrow.pack(fill="x", padx=SPACE_L, pady=(0, SPACE_S))

        # ── Event picker: prefill start/end from an existing annotation or
        # abnormal-event episode instead of requiring the user to know/type
        # the time by hand. Only shown when there is something to pick.
        _event_opts: "dict[str, tuple[float, float]]" = {}
        for _ann in self.analysis.annotations:
            _lbl = _ann.get("label") or "Annotation"
            _event_opts[f"📍 {_lbl}  [{_ann['t_start']:.0f}–{_ann['t_end']:.0f}s]"] = (
                float(_ann["t_start"]), float(_ann["t_end"]))
        for _ev in self.analysis.arrhythmia_events:
            _event_opts[f"⚠ {_ev.label}  [{_ev.t_start:.0f}–{_ev.t_end:.0f}s]"] = (
                float(_ev.t_start), float(_ev.t_end))

        def _on_event_pick(choice: str) -> None:
            times = _event_opts.get(choice)
            if times is None:
                return
            t0, t1 = times
            q_evt_start.delete(0, "end"); q_evt_start.insert(0, f"{t0:.1f}")
            q_evt_end.delete(0, "end");   q_evt_end.insert(0, f"{t1:.1f}")

        if _event_opts:
            ctk.CTkLabel(qrow, text="Event:", font=FONT_SMALL,
                         text_color=MUTED).pack(side="left")
            cb_event = ctk.CTkComboBox(
                qrow, width=190, height=28, font=FONT_LABEL,
                fg_color=BG, border_color=BORDER2, button_color=BORDER2,
                text_color=TEXT, dropdown_fg_color=BG, dropdown_text_color=TEXT,
                values=list(_event_opts.keys()), command=_on_event_pick)
            cb_event.set("Pick…")
            cb_event.pack(side="left", padx=(4, SPACE_M))

        ctk.CTkLabel(qrow, text="Event start (s):", font=FONT_SMALL,
                     text_color=MUTED).pack(side="left")
        q_evt_start = ctk.CTkEntry(qrow, width=70, height=28, font=FONT_LABEL,
                                   fg_color=BG, border_color=BORDER2, text_color=TEXT)
        if self.ui.nav_pos > 0:
            # Default to wherever the user is currently scrubbed to in the
            # Detection view -- most likely near the event they want to
            # compare around, since that's usually why they navigated there.
            q_evt_start.insert(0, str(int(self.ui.nav_pos)))
        q_evt_start.pack(side="left", padx=(4, SPACE_M))

        ctk.CTkLabel(qrow, text="Event end (s):", font=FONT_SMALL,
                     text_color=MUTED).pack(side="left")
        q_evt_end = ctk.CTkEntry(qrow, width=70, height=28, font=FONT_LABEL,
                                 fg_color=BG, border_color=BORDER2, text_color=TEXT,
                                 placeholder_text="= start")
        q_evt_end.pack(side="left", padx=(4, SPACE_M))

        ctk.CTkLabel(qrow, text="Window (min):", font=FONT_SMALL,
                     text_color=MUTED).pack(side="left")
        q_window = ctk.CTkEntry(qrow, width=55, height=28, font=FONT_LABEL,
                                fg_color=BG, border_color=BORDER2, text_color=TEXT)
        q_window.insert(0, "10")
        q_window.pack(side="left", padx=(4, SPACE_M))

        lbl_quick_status = ctk.CTkLabel(quick_frame, text="", font=FONT_MICRO,
                                        text_color=MUTED, anchor="w")

        def _apply_quick_fill() -> None:
            """Fill Segment A/B from an event time + a shared window duration.

            Segment A ("Before") = [event_start - window, event_start]
            Segment B ("After")  = [event_end,             event_end + window]
            event_end defaults to event_start (a point event) when left
            blank, so a single time + duration is enough for the common
            case, while still supporting a stimulus/injection that spans
            a start..end window of its own.
            """
            if not q_evt_start.get().strip():
                lbl_quick_status.configure(text="⚠  enter an event start time",
                                           text_color=RED)
                return
            try:
                evt_start = float(q_evt_start.get())
                evt_end_s = q_evt_end.get().strip()
                evt_end   = float(evt_end_s) if evt_end_s else evt_start
                window_m  = float(q_window.get())
                if window_m <= 0:
                    raise ValueError("window must be positive")
                if evt_end < evt_start:
                    raise ValueError("event end is before event start")
            except ValueError as exc:
                lbl_quick_status.configure(text=f"⚠  invalid input — {exc}",
                                           text_color=RED)
                return
            window_s = window_m * 60.0
            before_lo, before_hi = evt_start - window_s, evt_start
            after_lo,  after_hi  = evt_end,               evt_end + window_s
            # Clamp to the recording's own bounds -- on a short recording an
            # event near either edge used to produce a Before/After range
            # outside [0, dur] with nothing but a warning label; ▶ Compare
            # would then reject it outright as "outside recording".
            raw = (before_lo, before_hi, after_lo, after_hi)
            before_lo = max(0.0, min(before_lo, dur))
            before_hi = max(0.0, min(before_hi, dur))
            after_lo  = max(0.0, min(after_lo,  dur))
            after_hi  = max(0.0, min(after_hi,  dur))
            clamped = raw != (before_lo, before_hi, after_lo, after_hi)
            for entry, val in [(seg_entries[0]["lo"], before_lo),
                               (seg_entries[0]["hi"], before_hi),
                               (seg_entries[1]["lo"], after_lo),
                               (seg_entries[1]["hi"], after_hi)]:
                entry.delete(0, "end")
                entry.insert(0, f"{val:.1f}")
            for entry, txt in [(seg_entries[0]["lbl"], "Before"),
                               (seg_entries[1]["lbl"], "After")]:
                entry.delete(0, "end")
                entry.insert(0, txt)
            lbl_quick_status.configure(
                text=("⚠  clamped to recording bounds — " if clamped else "✓  ") +
                     f"Before [{before_lo:.0f}–{before_hi:.0f}s]  ·  "
                     f"After [{after_lo:.0f}–{after_hi:.0f}s]"
                     + (f"  (recording is {dur:.0f}s)" if clamped else ""),
                text_color=ORANGE if clamped else GREEN)

        ctk.CTkButton(qrow, text="Fill segments", width=110, height=28,
                     fg_color=TEAL, hover_color=TEAL_DARK, text_color="white",
                     font=FONT_BTN_SEC, corner_radius=6,
                     command=_apply_quick_fill).pack(side="left", padx=(SPACE_S, 0))
        lbl_quick_status.pack(padx=SPACE_L, pady=(0, SPACE_S), anchor="w", fill="x")

        ctrl = ctk.CTkFrame(win, fg_color="transparent")
        ctrl.pack(fill="x", padx=SPACE_L, pady=(0, SPACE_S))
        lbl_status = ctk.CTkLabel(ctrl, text="", font=FONT_SMALL, text_color=MUTED)
        lbl_status.pack(side="right", padx=(SPACE_M, 0))
        progress_bar = ctk.CTkProgressBar(ctrl, height=8, mode="indeterminate",
                                          progress_color=TEAL)
        progress_bar.pack(side="right", padx=(SPACE_M, 0))
        progress_bar.pack_forget()

        run_btn = ctk.CTkButton(ctrl, text="▶  Compare",
                                fg_color=TEAL, hover_color=TEAL_DARK,
                                text_color="white",
                                font=FONT_BTN_PRIMARY, height=34, corner_radius=8)
        run_btn.pack(side="left")
        # Length of the non-overlapping blocks each segment is summarised by
        # (compare_segment_blocks) -- the beat-level RR test this replaced
        # pseudo-replicated hundreds of autocorrelated beats.
        ctk.CTkLabel(ctrl, text="Block (s):", font=FONT_SMALL,
                     text_color=MUTED).pack(side="left", padx=(SPACE_L, 4))
        ent_block = ctk.CTkEntry(ctrl, width=55, height=28, font=FONT_LABEL,
                                 fg_color=BG, border_color=BORDER2, text_color=TEXT)
        ent_block.insert(0, f"{MouseECG.SEGMENT_BLOCK_S:g}")
        ent_block.pack(side="left")

        results_frame = ctk.CTkFrame(win, fg_color="transparent")
        results_frame.pack(fill="both", expand=True, padx=SPACE_L, pady=(0, SPACE_M))

        table_card = ctk.CTkFrame(results_frame, fg_color=PANEL, corner_radius=8)
        table_card.pack(side="left", fill="both", expand=True, padx=(0, SPACE_M))
        ctk.CTkLabel(table_card, text="Metric comparison",
                     font=FONT_SUBSECTION, text_color=MUTED,
                     anchor="w").pack(padx=SPACE_M, pady=(SPACE_M, SPACE_S), fill="x")
        lbl_mw = ctk.CTkLabel(table_card, text="", font=FONT_MICRO,
                              text_color=MUTED, anchor="w", justify="left",
                              wraplength=480)
        lbl_mw.pack(padx=SPACE_M, pady=(0, SPACE_XS), fill="x")
        tbl = ctk.CTkScrollableFrame(table_card, fg_color=BG, height=380,
                                     scrollbar_button_color=BORDER)
        tbl.pack(fill="both", expand=True, padx=SPACE_S, pady=(0, SPACE_M))
        for ci in range(4):
            tbl.grid_columnconfigure(ci, weight=1)
        # Placeholder row, cleared by _populate_table's winfo_children() sweep
        # the first time ▶ Compare actually runs.
        ctk.CTkLabel(tbl, text="Set segment ranges above, then click ▶ Compare",
                     font=FONT_SMALL, text_color=MUTED, anchor="w").grid(
            row=0, column=0, columnspan=4, sticky="ew", padx=SPACE_S, pady=SPACE_M)

        plot_card = ctk.CTkFrame(results_frame, fg_color=PANEL, corner_radius=8)
        plot_card.pack(side="right", fill="both", expand=True)
        ctk.CTkLabel(plot_card, text="RR tachograms (superimposed)",
                     font=FONT_SUBSECTION, text_color=MUTED,
                     anchor="w").pack(padx=SPACE_M, pady=(SPACE_M, SPACE_XS), fill="x")
        from ecg.ui.plots import CanvasSlot as _CS
        plot_slot = _CS(plot_card, 8, 4, toolbar=False, yscale_bar=True,
                         placeholder="Click ▶ Compare to overlay tachograms")

        _METRICS: "list[tuple[str, str, str]]" = list(self._SEGMENT_TABLE_METRICS)

        def _fmt(v: "Any") -> str:
            try:
                fv = float(v)
                if not np.isfinite(fv):
                    return "—"
                if abs(fv) >= 100: return f"{fv:.1f}"
                if abs(fv) >= 10:  return f"{fv:.2f}"
                return f"{fv:.3f}"
            except Exception:
                return str(v) if v else "—"

        def _populate_table(sa: dict, sb: dict) -> None:
            for w in tbl.winfo_children():
                w.destroy()
            for ci, (txt, col) in enumerate([
                ("Metric", TEXT), (sa["label"], _colors[0]),
                (sb["label"], _colors[1]), ("Δ%", MUTED),
            ]):
                ctk.CTkLabel(tbl, text=txt, font=FONT_BADGE,
                             text_color=col, anchor="w").grid(
                    row=0, column=ci, sticky="ew", padx=SPACE_S, pady=(SPACE_XS, SPACE_S))
            for ri, (key, name, unit) in enumerate(_METRICS, start=1):
                va = float(sa.get(key, float("nan")))
                vb = float(sb.get(key, float("nan")))
                bg = CARD if ri % 2 == 0 else BG
                lbl_txt = f"{name}  {unit}" if unit else name
                ctk.CTkLabel(tbl, text=lbl_txt, font=FONT_KPI_LABEL,
                             text_color=MUTED, fg_color=bg,
                             anchor="w").grid(row=ri, column=0, sticky="ew", padx=SPACE_S, pady=SPACE_XS)
                ctk.CTkLabel(tbl, text=_fmt(va), font=FONT_KPI_LABEL,
                             text_color=TEXT, fg_color=bg,
                             anchor="e").grid(row=ri, column=1, sticky="ew", padx=SPACE_S, pady=SPACE_XS)
                ctk.CTkLabel(tbl, text=_fmt(vb), font=FONT_KPI_LABEL,
                             text_color=TEXT, fg_color=bg,
                             anchor="e").grid(row=ri, column=2, sticky="ew", padx=SPACE_S, pady=SPACE_XS)
                if np.isfinite(va) and np.isfinite(vb) and abs(va) > 1e-9:
                    delta = 100.0 * (vb - va) / abs(va)
                    ctk.CTkLabel(tbl, text=f"{'▲' if delta>0 else '▼'} {abs(delta):.1f}%",
                                 font=FONT_MICRO, text_color=MUTED, fg_color=bg,
                                 anchor="e").grid(row=ri, column=3, sticky="ew", padx=SPACE_S, pady=SPACE_XS)

        def _populate_plot(sa: dict, sb: dict) -> None:
            _t_a, _r_a, _la, _ca = sa["t_rr"], sa["rr_ms"], sa["label"], _colors[0]
            _t_b, _r_b, _lb, _cb = sb["t_rr"], sb["rr_ms"], sb["label"], _colors[1]
            def _draw(fig):
                ax = fig.add_subplot(111)
                style_axes(ax)
                if len(_t_a) and len(_r_a):
                    ax.plot(_t_a - _t_a[0], _r_a, color=_ca, lw=1.0, alpha=0.85, label=_la)
                    ax.axhline(float(np.mean(_r_a)), color=_ca, lw=1.0, ls="--", alpha=0.5)
                if len(_t_b) and len(_r_b):
                    ax.plot(_t_b - _t_b[0], _r_b, color=_cb, lw=1.0, alpha=0.85, label=_lb)
                    ax.axhline(float(np.mean(_r_b)), color=_cb, lw=1.0, ls="--", alpha=0.5)
                ax.set_xlabel("Time within segment (s)", fontsize=9)
                ax.set_ylabel("RR (ms)", fontsize=9)
                ax.set_title("RR intervals — segments superimposed", loc="left", fontsize=9)
                ax.legend(framealpha=0, fontsize=8)
            plot_slot.update(_draw)

        _last_cmp: "dict[str, Any]" = {"cmp": None}

        def _run_compare() -> None:
            try:
                lo_a = float(seg_entries[0]["lo"].get()); hi_a = float(seg_entries[0]["hi"].get())
                la   = seg_entries[0]["lbl"].get().strip() or "Segment A"
                lo_b = float(seg_entries[1]["lo"].get()); hi_b = float(seg_entries[1]["hi"].get())
                lb   = seg_entries[1]["lbl"].get().strip() or "Segment B"
            except ValueError:
                lbl_status.configure(text="⚠  Invalid time values", text_color=RED)
                return
            try:
                block_s = float(ent_block.get())
                if not (np.isfinite(block_s) and block_s > 0):
                    raise ValueError("block length")
            except ValueError:
                lbl_status.configure(text="⚠  Block length must be a positive number of seconds",
                                     text_color=RED)
                return
            for lo, hi, lbl_s in [(lo_a, hi_a, "A"), (lo_b, hi_b, "B")]:
                if lo >= hi:
                    lbl_status.configure(text=f"⚠  Segment {lbl_s}: start ≥ end", text_color=RED)
                    return
                if lo < 0 or hi > dur:
                    lbl_status.configure(
                        text=f"⚠  Segment {lbl_s}: outside recording ({dur:.1f} s)",
                        text_color=RED)
                    return
            run_btn.configure(state="disabled")
            progress_bar.pack(side="right", padx=(SPACE_M, 0))
            progress_bar.start()
            lbl_status.configure(text="  Computing…", text_color=ORANGE)
            # HRV inputs read on the main thread: the user's RR bounds and
            # the beats removed by artifact correction (excluded spans).
            rr_lo, rr_hi = self.analysis_ctrl.hrv_rr_bounds()
            seg_breaks   = self.analysis_ctrl.hrv_breaks(rp)

            import threading as _th
            def _worker():
                # Own try/except -- anything raising outside
                # compute_segment_stats()' own handling used to leave the
                # dialog stuck on "Computing…" with Compare still disabled.
                try:
                    from ecg.core.analysis import compare_segments, compute_segment_stats
                    seg_kw = dict(lf_band=lf, hf_band=hf, vlf_band=vlf,
                                  rr_min_ms=rr_lo, rr_max_ms=rr_hi,
                                  breaks=seg_breaks, block_s=block_s)
                    sa = compute_segment_stats(sig, rp, fs, lo_a, hi_a, la, **seg_kw)
                    sb = compute_segment_stats(sig, rp, fs, lo_b, hi_b, lb, **seg_kw)
                    # Descriptive block-mean comparison (no beat-level test:
                    # beats of one recording are not independent replicates).
                    cmp = compare_segments(sa, sb)
                except Exception as exc:
                    log.exception("Segment comparison worker failed")
                    win.after(0, lambda e=exc: _on_worker_error(e))
                    return
                win.after(0, lambda: _on_done(sa, sb, cmp))

            def _on_worker_error(exc: Exception) -> None:
                run_btn.configure(state="normal")
                progress_bar.stop(); progress_bar.pack_forget()
                lbl_status.configure(text=f"⚠  Comparison failed — {exc}", text_color=RED)

            def _on_done(sa: dict, sb: dict, cmp: dict) -> None:
                run_btn.configure(state="normal")
                progress_bar.stop(); progress_bar.pack_forget()
                # Persist for Prism export
                self._last_seg_a = sa
                self._last_seg_b = sb
                _last_cmp["cmp"] = cmp
                errs = [s["error"] for s in (sa, sb) if s.get("error")]
                if errs:
                    lbl_status.configure(text=f"⚠  {errs[0]}", text_color=RED)
                else:
                    # LF/HF withheld where excluded RR intervals are too long
                    # to interpolate across (compute_segment_stats).
                    withheld = [s["label"] for s in (sa, sb) if s.get("freq_withheld_reason")]
                    lbl_status.configure(
                        text=f"✓  {sa['label']}: {sa['n_beats']} beats  ·  "
                             f"{sb['label']}: {sb['n_beats']} beats"
                             + (f"  ·  LF/HF withheld for {', '.join(withheld)} "
                                f"(long RR gaps)" if withheld else ""),
                        text_color=ORANGE if withheld else GREEN)
                    self._set_btn_enabled(export_btn, True, GREEN_DARK, GREEN)   # unlock Export button
                from ecg.core.analysis import segment_comparison_text
                lbl_mw.configure(text=segment_comparison_text(cmp, sa["label"], sb["label"]))
                _populate_table(sa, sb)
                _populate_plot(sa, sb)

            _th.Thread(target=_worker, daemon=True).start()

        # ── Export function ───────────────────────────────────────────────
        _METRICS_EXPORT: "list[tuple[str, str, str]]" = list(self._SEGMENT_TABLE_METRICS)

        def _export_comparison() -> None:
            sa = self._last_seg_a
            sb = self._last_seg_b
            if sa is None or sb is None:
                messagebox.showwarning("No results",
                                       "Run the comparison first.", parent=win)
                return

            sub  = self.ent_subject.get().strip() if self.ent_subject else "subject"
            ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = f"{sub}_comparison_{ts}"

            path = filedialog.asksaveasfilename(
                parent=win,
                title="Export segment comparison",
                defaultextension=".xlsx",
                initialfile=f"{default_name}.xlsx",
                filetypes=[
                    ("Excel workbook", "*.xlsx"),
                    ("CSV",            "*.csv"),
                    ("All files",      "*.*"),
                ],
            )
            if not path:
                return

            la = sa["label"]
            lb = sb["label"]

            # ── Build metric comparison DataFrame ─────────────────────────
            rows = []
            for key, name, unit in _METRICS_EXPORT:
                va = sa.get(key, float("nan"))
                vb = sb.get(key, float("nan"))
                try:
                    va = float(va) if va is not None else float("nan")
                    vb = float(vb) if vb is not None else float("nan")
                except (TypeError, ValueError):
                    va = vb = float("nan")
                delta = (100.0 * (vb - va) / abs(va)
                         if np.isfinite(va) and np.isfinite(vb) and abs(va) > 1e-9
                         else float("nan"))
                rows.append({
                    "Metric":      f"{name}  ({unit})" if unit else name,
                    la:            va,
                    lb:            vb,
                    "Δ% (A→B)":    delta,
                })
            cmp = _last_cmp.get("cmp")
            if cmp is not None:
                # The descriptive-only caveat travels with the numbers.
                from ecg.core.analysis import segment_comparison_text
                rows.append({"Metric": "Note: " + segment_comparison_text(
                                 cmp, la, lb).replace("\n", "  "),
                             la: None, lb: None, "Δ% (A→B)": None})
            metrics_df = pd.DataFrame(rows)

            # ── RR series per segment ─────────────────────────────────────
            rr_a = pd.DataFrame({
                "Time (s)": np.asarray(sa["t_rr"], dtype=float),
                "RR (ms)":  np.asarray(sa["rr_ms"], dtype=float),
            })
            rr_b = pd.DataFrame({
                "Time (s)": np.asarray(sb["t_rr"], dtype=float),
                "RR (ms)":  np.asarray(sb["rr_ms"], dtype=float),
            })

            # ── Write ─────────────────────────────────────────────────────
            if path.lower().endswith(".csv"):
                metrics_df.to_csv(path, index=False, float_format="%.4f")
                lbl_status.configure(
                    text=f"✓  CSV saved → {os.path.basename(path)}",
                    text_color=GREEN)
            else:
                try:
                    import openpyxl
                    from openpyxl.styles import Font
                    from openpyxl.utils import get_column_letter
                except ImportError:
                    messagebox.showerror("Missing dependency",
                                         "openpyxl required: pip install openpyxl",
                                         parent=win)
                    return

                wb = openpyxl.Workbook()

                # ── Sheet 1: Metric comparison ────────────────────────────
                ws = wb.active
                assert ws is not None
                ws.title = "Comparison"

                headers = list(metrics_df.columns)
                for ci, h in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=ci, value=h)
                    cell.font = Font(bold=True)

                for ri, row_d in enumerate(metrics_df.itertuples(index=False), start=2):
                    for ci, val in enumerate(row_d, start=1):
                        ws.cell(row=ri, column=ci, value=val)

                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for col in ws.columns:
                    if not col:
                        continue
                    first_cell = col[0]
                    if first_cell.column is None:
                        continue
                    max_len = max(
                        (len(str(c.value)) for c in col if c.value), default=8)
                    ws.column_dimensions[
                        get_column_letter(first_cell.column)].width = min(
                        max_len + 3, 30)

                # ── Sheet 2: RR series ────────────────────────────────────
                def _rr_sheet(name: str, rr_df: pd.DataFrame) -> None:
                    ws2 = wb.create_sheet(name)
                    for ci, col in enumerate(rr_df.columns, start=1):
                        c = ws2.cell(row=1, column=ci, value=col)
                        c.font = Font(bold=True)
                    for ri, row_vals in enumerate(rr_df.itertuples(index=False),
                                                   start=2):
                        for ci, v in enumerate(row_vals, start=1):
                            ws2.cell(row=ri, column=ci, value=round(float(v), 4))
                    ws2.freeze_panes = "A2"
                    ws2.auto_filter.ref = ws2.dimensions
                    ws2.column_dimensions["A"].width = 14
                    ws2.column_dimensions["B"].width = 12

                _rr_sheet(f"RR — {la}", rr_a)
                _rr_sheet(f"RR — {lb}", rr_b)

                # ── Sheet: per-block means (the comparison's data) ────────
                ws_b = wb.create_sheet("Blocks")
                for ci, h in enumerate(["Segment", "Block start (s)", "N RR",
                                        "Mean RR (ms)", "HR (bpm, 60000/mean NN)"],
                                       start=1):
                    ws_b.cell(row=1, column=ci, value=h).font = Font(bold=True)
                rb = 2
                for seg in (sa, sb):
                    blk = seg.get("blocks") or {}
                    for t0, n_rr, rr_m, hr_m in zip(blk.get("t_start", []), blk.get("n_rr", []),
                                                    blk.get("rr_mean", []), blk.get("hr_mean", [])):
                        for ci, v in enumerate([seg["label"], round(float(t0), 2), int(n_rr),
                                                round(float(rr_m), 3), round(float(hr_m), 2)],
                                               start=1):
                            ws_b.cell(row=rb, column=ci, value=v)
                        rb += 1
                ws_b.freeze_panes = "A2"
                for col_l, w in zip("ABCDE", (16, 16, 8, 14, 24)):
                    ws_b.column_dimensions[col_l].width = w

                # ── Sheet 3: Summary stats ────────────────────────────────
                ws3 = wb.create_sheet("Summary")
                ws3.cell(row=1, column=1, value="Subject").font = Font(bold=True)
                ws3.cell(row=1, column=2, value=sub)
                ws3.cell(row=2, column=1, value="Export date").font = Font(bold=True)
                ws3.cell(row=2, column=2,
                         value=datetime.now().strftime("%Y-%m-%d %H:%M"))
                ws3.cell(row=3, column=1, value="LF band (Hz)").font = Font(bold=True)
                ws3.cell(row=3, column=2, value=f"{lf[0]:.2f}–{lf[1]:.2f}")
                ws3.cell(row=4, column=1, value="HF band (Hz)").font = Font(bold=True)
                ws3.cell(row=4, column=2, value=f"{hf[0]:.2f}–{hf[1]:.2f}")
                if cmp is not None:
                    hr_c = cmp["metrics"]["hr_mean"]
                    p_d  = hr_c.get("p_descriptive")
                    for ri, (k, v) in enumerate([
                        ("Block length (s)", cmp["block_s"]),
                        (f"Blocks used — {la}", cmp["n_blocks_a"]),
                        (f"Blocks used — {lb}", cmp["n_blocks_b"]),
                        ("Blocks dropped (< 50 % valid RR)",
                         f"{cmp['n_dropped_a']} / {cmp['n_dropped_b']}"),
                        ("HR block means Δ (B − A, bpm)", hr_c["diff"]),
                        ("Descriptive Mann-Whitney p (block means)",
                         p_d if p_d is not None else "n/a (< 3 blocks)"),
                        ("Note", cmp["note"]),
                    ], start=5):
                        ws3.cell(row=ri, column=1, value=k).font = Font(bold=True)
                        ws3.cell(row=ri, column=2, value=v)
                for col in ws3.columns:
                    ws3.column_dimensions[
                        get_column_letter(col[0].column)].width = 22

                wb.save(path)
                lbl_status.configure(
                    text=f"✓  Exported → {os.path.basename(path)}",
                    text_color=GREEN)

            # ── Also save the plot as PNG next to the xlsx ─────────────────
            png_path = os.path.splitext(path)[0] + "_plot.png"
            try:
                export_fig = Figure(
                    figsize=(10, 4), dpi=200,
                    facecolor=PLOT.get("bg", "#1A1A2E"), layout="constrained")
                if hasattr(export_fig, "set_constrained_layout_pads"):
                    # Was previously called with left=/right=/top=/bottom= --
                    # not real parameters of this method (only w_pad, h_pad,
                    # wspace, hspace are), so this raised TypeError on every
                    # single call, silently caught by the except block below
                    # at DEBUG level: this companion PNG has never actually
                    # been written. Reuse CanvasSlot's own pad constants
                    # instead of a third hardcoded copy of the numbers.
                    getattr(export_fig, "set_constrained_layout_pads")(
                        w_pad=CanvasSlot._CL_PAD, h_pad=CanvasSlot._CL_PAD,
                        wspace=CanvasSlot._CL_SPACE, hspace=CanvasSlot._CL_SPACE)
                if plot_slot._draw_fn is not None:
                    plot_slot._draw_fn(export_fig)
                    export_fig.savefig(
                        png_path, dpi=200,
                        facecolor=export_fig.get_facecolor(),
                        bbox_inches="tight")
                    plt.close(export_fig)
                    lbl_status.configure(
                        text=lbl_status.cget("text") + f"  +  {os.path.basename(png_path)}",
                        text_color=GREEN)
            except Exception as exc:
                log.debug("comparison PNG export: %s", exc)

        run_btn.configure(command=_run_compare)

        # Export button — disabled until comparison has run. Built with an
        # explicit muted fg_color (see _set_btn_enabled) rather than relying
        # on CTkButton's default disabled treatment, which only dims the
        # text and left this looking clickable while genuinely disabled.
        export_btn = ctk.CTkButton(
            ctrl, text="📊  Export",
            command=_export_comparison,
            fg_color=BORDER, hover_color=BORDER2, text_color=MUTED,
            font=FONT_BTN_PRIMARY, height=34, corner_radius=8,
            state="disabled",
        )
        export_btn.pack(side="left", padx=(SPACE_M, 0))

        ctk.CTkLabel(win,
                     text=f"Tip: equal durations give a fair comparison  ·  "
                          f"LF {lf[0]:.2f}–{lf[1]:.2f} Hz  HF {hf[0]:.2f}–{hf[1]:.2f} Hz",
                     font=FONT_MICRO, text_color=LIGHT, anchor="w").pack(
            fill="x", padx=SPACE_L, pady=(0, SPACE_S))
